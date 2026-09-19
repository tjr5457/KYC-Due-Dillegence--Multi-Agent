from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

def emu_to_inch(emu):
    return round(emu / 914400, 2)

def iter_shapes(shapes, level=0):
    items = []
    for shape in shapes:
        info = {
            "name": shape.name,
            "left": emu_to_inch(shape.left) if shape.left is not None else None,
            "top": emu_to_inch(shape.top) if shape.top is not None else None,
            "width": emu_to_inch(shape.width) if shape.width is not None else None,
            "height": emu_to_inch(shape.height) if shape.height is not None else None,
            "text": shape.text.strip()[:60] if shape.has_text_frame else "",
        }
        items.append(info)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            items.extend(iter_shapes(shape.shapes, level+1))
    return items

def analyze_ppt(path):
    prs = Presentation(path)
    W = emu_to_inch(prs.slide_width)
    H = emu_to_inch(prs.slide_height)
    print(f"Slide size: {W}\" x {H}\"")
    print(f"Total slides: {len(prs.slides)}\n")

    slide_titles = []
    slide_margins = []

    for i, slide in enumerate(prs.slides):
        shapes = iter_shapes(slide.shapes)
        
        # Find leftmost and rightmost content positions
        lefts = [s["left"] for s in shapes if s["left"] is not None and s["left"] > 0]
        rights = [s["left"] + s["width"] for s in shapes if s["left"] is not None and s["width"] is not None]
        tops = [s["top"] for s in shapes if s["top"] is not None and s["top"] > 0.05]
        
        min_left  = min(lefts) if lefts else None
        max_right = max(rights) if rights else None
        min_top   = min(tops) if tops else None
        
        # Get title text
        title_text = ""
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                title_text = shape.text.strip()[:50]
                break
        
        slide_titles.append(title_text)
        slide_margins.append({
            "slide": i+1,
            "title": title_text,
            "left_margin": min_left,
            "right_edge": max_right,
            "top_margin": min_top,
        })
        
        print(f"--- Slide {i+1}: '{title_text}' ---")
        print(f"  Left margin : {min_left}\"")
        print(f"  Right edge  : {max_right}\" (slide width: {W}\")")
        print(f"  Top margin  : {min_top}\"")
        print()

    # Check consistency
    print("=" * 60)
    print("ALIGNMENT CONSISTENCY REPORT")
    print("=" * 60)
    
    left_margins = [s["left_margin"] for s in slide_margins if s["left_margin"] is not None]
    right_edges  = [s["right_edge"] for s in slide_margins if s["right_edge"] is not None]
    top_margins  = [s["top_margin"] for s in slide_margins if s["top_margin"] is not None]
    
    common_left  = round(sum(left_margins)/len(left_margins), 2) if left_margins else None
    common_right = round(sum(right_edges)/len(right_edges), 2) if right_edges else None

    issues_found = False
    for s in slide_margins:
        problems = []
        if s["left_margin"] is not None and abs(s["left_margin"] - common_left) > 0.3:
            problems.append(f"LEFT margin {s['left_margin']}\" (avg={common_left}\")")
        if s["right_edge"] is not None and abs(s["right_edge"] - common_right) > 0.5:
            problems.append(f"RIGHT edge {s['right_edge']}\" (avg={common_right}\")")
        
        if problems:
            issues_found = True
            print(f"[MISALIGNED] Slide {s['slide']} '{s['title']}': {', '.join(problems)}")
        else:
            print(f"[OK]         Slide {s['slide']} '{s['title']}'")
    
    if not issues_found:
        print("\nAll slides are consistently aligned!")

if __name__ == "__main__":
    analyze_ppt("updated ppt.pptx")
