"""
Updates team member names and roles in TCS-AMD-KYC-8Slides.pptx
"""
from pptx import Presentation

TEAM_TEXT = (
    "Juvviraju Tarra       : AI Architect – Multi-Agent Design, LangGraph Pipeline\n"
    "Kodanda Rama Sai Divvi: Developer   – KYC Logic, API & AMD ROCm Integration\n"
    "Dhinesh N             : Full-Stack Developer – Streamlit Dashboard, FastAPI Integration,\n"
    "                        Real-time Pipeline Monitoring & Business Case"
)

prs = Presentation("TCS-AMD-KYC-8Slides.pptx")
slides = list(prs.slides)
s2 = slides[1]  # Slide 2 – Basic Information

updated = False
for shape in s2.shapes:
    if shape.has_text_frame:
        text = shape.text.strip()
        # Target the member names/roles shape
        if "[Member" in text or "Member 1" in text or "AI/ML" in text or "Juvvi" in text:
            tf = shape.text_frame
            tf.word_wrap = True
            # Clear all paragraphs beyond the first
            while len(tf.paragraphs) > 1:
                p = tf.paragraphs[-1]._p
                p.getparent().remove(p)
            para = tf.paragraphs[0]
            while len(para.runs) > 1:
                r = para.runs[-1]._r
                r.getparent().remove(r)
            if para.runs:
                para.runs[0].text = TEAM_TEXT
            else:
                from pptx.oxml.ns import qn
                from lxml import etree
                r_elem = etree.SubElement(para._p, qn("a:r"))
                etree.SubElement(r_elem, qn("a:rPr"), attrib={"lang": "en-US", "dirty": "0"})
                t_elem = etree.SubElement(r_elem, qn("a:t"))
                t_elem.text = TEAM_TEXT
            print(f"Updated shape: {repr(shape.name)}")
            updated = True
            break

if not updated:
    # Fallback: print all shapes to debug
    print("Shape not found. Dumping slide 2 shapes:")
    for shape in s2.shapes:
        if shape.has_text_frame:
            print(f"  {repr(shape.name)}: {repr(shape.text[:100])}")
else:
    prs.save("TCS-AMD-KYC-8Slides.pptx")
    print("Saved successfully!")
